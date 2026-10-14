"""
documents.py — turn simple Markdown into a PDF, Word (.docx) or Excel (.xlsx) file.

The model only has to write plain text with a few marks; this code does the layout, so the result does not
depend on a small model getting reportlab or python-docx code right.

Supported marks:   # / ## / ### headings     - or * bullets     1. numbered lists
                   **bold**                   | a | b |  tables (first row is the header)     --- ignored
"""

import os
import re
from xml.sax.saxutils import escape

# Text like [Your Name] or [Company Address] means the document was written without the real details.
PLACEHOLDER = re.compile(r"\[(?:your|company|date|hiring|recipient|insert|name|address|phone|email)[^\]\n]{0,40}\]", re.I)
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def parse(text):
    """Markdown-ish text -> list of blocks: ("h", level, text) ("p", text) ("ul", [..]) ("ol", [..]) ("table", [[..]])."""
    blocks, para, lines = [], [], str(text).replace("\r\n", "\n").split("\n")

    def flush():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            flush()
        elif re.match(r"^-{3,}$|^\*{3,}$|^_{3,}$", s):
            flush()
        elif re.match(r"^#{1,4}\s+\S", s):
            flush()
            level = len(s) - len(s.lstrip("#"))
            blocks.append(("h", min(level, 3), s.lstrip("#").strip()))
        elif re.match(r"^[-*•]\s+\S", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*•]\s+\S", lines[i]):
                items.append(re.sub(r"^\s*[-*•]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"^\d+[.)]\s+\S", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+\S", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
            continue
        elif s.startswith("|") and s.count("|") >= 2:
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not TABLE_SEP.match(row):
                    rows.append([c.strip() for c in row.strip("|").split("|")])
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                blocks.append(("table", [r + [""] * (width - len(r)) for r in rows]))
            continue
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


def _plain(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", str(text))


# ── PDF ──────────────────────────────────────────────────────────────────
def _fonts():
    """Use Arial if this Mac has it (covers more characters), else the built-in Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    base = "/System/Library/Fonts/Supplemental/"
    files = {"Doc": "Arial.ttf", "Doc-Bold": "Arial Bold.ttf", "Doc-Italic": "Arial Italic.ttf",
             "Doc-BoldItalic": "Arial Bold Italic.ttf"}
    try:
        if "Doc" not in pdfmetrics.getRegisteredFontNames():
            for name, f in files.items():
                pdfmetrics.registerFont(TTFont(name, base + f))
            pdfmetrics.registerFontFamily("Doc", normal="Doc", bold="Doc-Bold", italic="Doc-Italic", boldItalic="Doc-BoldItalic")
        return "Doc", "Doc-Bold", True
    except Exception:
        return "Helvetica", "Helvetica-Bold", False


def _pdf_text(text, unicode_ok):
    s = escape(str(text))
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    if not unicode_ok:
        s = s.encode("cp1252", "replace").decode("cp1252")
    return s


def build_pdf(path, title, blocks):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold, uni = _fonts()
    body = ParagraphStyle("body", fontName=reg, fontSize=10.5, leading=15, spaceAfter=6)
    heads = {1: ParagraphStyle("h1", parent=body, fontName=bold, fontSize=16, leading=20, spaceBefore=10, spaceAfter=6),
             2: ParagraphStyle("h2", parent=body, fontName=bold, fontSize=13, leading=17, spaceBefore=8, spaceAfter=4),
             3: ParagraphStyle("h3", parent=body, fontName=bold, fontSize=11, leading=15, spaceBefore=6, spaceAfter=3)}
    cell = ParagraphStyle("cell", parent=body, fontSize=9.5, leading=12, spaceAfter=0)
    cellb = ParagraphStyle("cellb", parent=cell, fontName=bold)
    title_style = ParagraphStyle("title", parent=body, fontName=bold, fontSize=22, leading=27, spaceAfter=10)

    flow = []
    if title:
        flow.append(Paragraph(_pdf_text(title, uni), title_style))
    for b in blocks:
        kind = b[0]
        if kind == "h":
            flow.append(Paragraph(_pdf_text(b[2], uni), heads[b[1]]))
        elif kind == "p":
            flow.append(Paragraph(_pdf_text(b[1], uni), body))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(_pdf_text(t, uni), body), leftIndent=14) for t in b[1]]
            flow.append(ListFlowable(items, bulletType="bullet" if kind == "ul" else "1", start="•" if kind == "ul" else 1,
                                     leftIndent=14, bulletFontName=reg))
            flow.append(Spacer(1, 3))
        elif kind == "table":
            rows = [[Paragraph(_pdf_text(c, uni), cellb if r == 0 else cell) for c in row] for r, row in enumerate(b[1])]
            t = Table(rows, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9e4d6")),
                                   ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8b2a2")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            flow += [t, Spacer(1, 8)]
    if not flow:
        flow.append(Paragraph(" ", body))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title or os.path.splitext(os.path.basename(path))[0])
    doc.build(flow)
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "%d page%s" % (len(pdf.pages), "" if len(pdf.pages) == 1 else "s")
    except Exception:
        return "PDF"


# ── Word ─────────────────────────────────────────────────────────────────
def _runs(paragraph, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if part:
            paragraph.add_run(part).bold = bool(i % 2)


def build_docx(path, title, blocks):
    import docx
    d = docx.Document()
    if title:
        d.add_heading(title, level=0)
    for b in blocks:
        kind = b[0]
        if kind == "h":
            d.add_heading(_plain(b[2]), level=b[1])
        elif kind == "p":
            _runs(d.add_paragraph(), b[1])
        elif kind in ("ul", "ol"):
            for t in b[1]:
                _runs(d.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), t)
        elif kind == "table":
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    p = t.cell(r, c).paragraphs[0]
                    _runs(p, val)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
    d.save(path)
    return "Word document"


# ── Excel ────────────────────────────────────────────────────────────────
def _cell_value(v):
    s = _plain(v).strip()
    num = s.replace(",", "")
    try:
        if re.fullmatch(r"-?\d+", num):
            return int(num)
        if re.fullmatch(r"-?\d*\.\d+", num):
            return float(num)
    except ValueError:
        pass
    return s


def build_xlsx(path, title, blocks):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    tables = [b[1] for b in blocks if b[0] == "table"]
    for n, rows in enumerate(tables, 1):
        ws = wb.create_sheet("Sheet%d" % n if len(tables) > 1 else "Data")
        for r, row in enumerate(rows, 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c)
                cell.value = _cell_value(val) if r > 1 else _plain(val).strip()
                if isinstance(cell.value, str):
                    cell.data_type = "s"          # text that starts with "=" must stay text, not become a formula
                if r == 1:
                    cell.font = Font(bold=True)
        for c in range(1, len(rows[0]) + 1):
            width = max(len(_plain(str(row[c - 1]))) for row in rows)
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 60)
    others = [b for b in blocks if b[0] != "table"]
    if others or not tables:
        ws = wb.create_sheet("Notes")
        r = 1
        if title:
            ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
            r += 2
        for b in others:
            items = b[1] if b[0] in ("ul", "ol") else [b[2] if b[0] == "h" else b[1]]
            for t in items:
                cell = ws.cell(row=r, column=1, value=_plain(t))
                cell.data_type = "s"
                if b[0] == "h":
                    cell.font = Font(bold=True)
                r += 1
        ws.column_dimensions["A"].width = 100
    wb.save(path)
    return "%d sheet%s" % (len(wb.sheetnames), "" if len(wb.sheetnames) == 1 else "s")


BUILDERS = {".pdf": build_pdf, ".docx": build_docx, ".xlsx": build_xlsx}
